# -*- coding: utf-8 -*-
"""Extract-as-is rebuild of the PDF report into Word (.docx).

Two variants are produced:
  * "faithful" – keeps the asterisks exactly where the original PDF has them
  * "clean"    – same text without the asterisk artifacts
Both use the original typeface (B Nazanin), full RTL layout and the
decorative page frame of the source document.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CS_FONT = "B Nazanin"       # font embedded in the source PDF
LATIN_FONT = "Times New Roman"

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
    "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
    "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz",
    "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
    "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]
SECT_ORDER = [
    "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc", "pgBorders",
    "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
    "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings",
    "sectPrChange",
]


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _reorder(el, order):
    rank = {qn("w:" + name): i for i, name in enumerate(order)}
    children = list(el)
    children.sort(key=lambda c: rank.get(c.tag, len(order)))
    for c in children:
        el.remove(c)
    for c in children:
        el.append(c)


def _rtl_paragraph(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))


def _style_run(run, size=13, bold=False):
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), LATIN_FONT)
    rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    rFonts.set(qn("w:cs"), CS_FONT)
    rPr.append(rFonts)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)
    rPr.append(OxmlElement("w:rtl"))
    if bold:
        rPr.append(OxmlElement("w:bCs"))


def add_par(doc, text="", size=13, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            space_after=8, space_before=0, line=1.55):
    p = doc.add_paragraph()
    _rtl_paragraph(p)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if text:
        _style_run(p.add_run(text), size=size, bold=bold)
    return p


def add_item(doc, number, lead, body, star_mark=False, size=13):
    """Numbered bullet: bold lead phrase followed by the explanatory text."""
    p = doc.add_paragraph()
    _rtl_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.55
    head = f"{number}*. {lead}" if star_mark else f"{number}. {lead}"
    _style_run(p.add_run(head), size=size, bold=True)
    _style_run(p.add_run(" " + body), size=size, bold=False)
    return p


def add_page_border(section):
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "double")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), "1F3864")
        borders.append(el)
    section._sectPr.append(borders)


# --------------------------------------------------------------------------- #
# document content (transcribed line-by-line from the PDF)
# --------------------------------------------------------------------------- #
TITLE = ("گزارش عملکرد برنامه تلویزیونی «به خانه برمی‌گردیم» شبکه تهران سیما "
         "در قالب طرح «ایران‌جان» ـ استان مرکزی")
SUBJECT = ("موضوع: گزارش حضور و تولید و پخش برنامه با محوریت معرفی ظرفیت‌های "
           "استان مرکزی در قالب طرح ایران‌جان")

INTRO_1 = ("در راستای تحقق اهداف طرح ملی «ایران‌جان» و با هدف معرفی ظرفیت‌های "
           "فرهنگی، هنری، تاریخی، گردشگری و انسانی استان‌های کشور، برنامه "
           "تلویزیونی «به خانه برمی‌گردیم» شبکه تهران سیما به مدت دو روز در "
           "استان مرکزی حضور یافت و تمامی بخش‌های برنامه در این دو روز به معرفی "
           "و بازتاب ظرفیت‌ها و توانمندی‌های این استان اختصاص پیدا کرد.")
INTRO_2 = ("این حضور با رویکرد معرفی هویت بومی و محلی استان مرکزی، انعکاس "
           "جلوه‌های فرهنگ و هنر ایرانی، توجه به مفاخر و نخبگان منطقه و همچنین "
           "معرفی جاذبه‌های تاریخی و گردشگری آن برنامه‌ریزی و اجرا شد.")
LEAD_IN = ("مهم‌ترین دستاوردهای حضور برنامه «به خانه برمی‌گردیم» در استان مرکزی "
           "را می‌توان در محورهای زیر جمع‌بندی کرد:")

ITEMS = [
    ("۱", "اختصاص کامل برنامه به طرح ایران‌جان:",
     "تمامی محتوای برنامه در هر دو روز حضور در استان مرکزی، در راستای معرفی این "
     "استان و بازتاب ظرفیت‌های آن تنظیم و اجرا شد."),
    ("۲", "اختصاص ۱۵۰ دقیقه آنتن تلویزیونی:",
     "در مجموع دو روز، ۱۵۰ دقیقه از آنتن برنامه با محوریت استان مرکزی و معرفی "
     "ظرفیت‌های آن همراه بود."),
    ("۳", "بهره‌گیری از ظرفیت‌های بومی و محلی:",
     "حضور آشپزان، مفاخر، نخبگان، هنرمندان و فعالان گردشگری، زمینه معرفی مستقیم "
     "توانمندی‌های انسانی و فرهنگی استان را فراهم کرد."),
    ("۴", "معرفی هویت تاریخی و معماری منطقه:",
     "انتخاب خانه‌های تاریخی و بافت‌های قدیمی به‌عنوان لوکیشن‌های برنامه، به "
     "تقویت جنبه‌های بصری و فرهنگی تولید کمک کرد."),
    ("۵", "نمایش هنرهای بومی در بستر واقعی فعالیت:",
     "حضور هنرمندان و فعالان صنایع‌دستی در فضاهای مرتبط با کار آنان، امکان "
     "بازنمایی عینی هنرهای سنتی و مهارت‌های بومی را فراهم ساخت."),
    ("۶", "تنوع‌بخشی به ساختار برنامه:",
     "ترکیب گفت‌وگوها، حضور مهمانان، لوکیشن‌های متنوع و پلی‌بک‌های مرتبط، به "
     "ایجاد ساختاری چندوجهی برای معرفی استان کمک کرد."),
    ("۷", "تقویت نقش رسانه در معرفی ظرفیت‌های استانی:",
     "این حضور، نمونه‌ای از به‌کارگیری ظرفیت یک برنامه تلویزیونی روزانه برای "
     "توجه هدفمند به هویت فرهنگی و اجتماعی یکی از استان‌های کشور بود."),
]

CONC_1 = ("حضور دو روزه برنامه «به خانه برمی‌گردیم» شبکه تهران سیما در استان "
          "مرکزی، در چارچوب طرح ملی «ایران‌جان»، با اختصاص کامل محتوای برنامه به "
          "معرفی ظرفیت‌های این استان و بهره‌گیری از ۱۵۰ دقیقه آنتن تلویزیونی "
          "انجام شد.")
CONC_2 = ("در این مدت، تلاش شد با استفاده از ترکیب مهمانان بومی و محلی، معرفی "
          "مفاخر و نخبگان، حضور هنرمندان و فعالان صنایع‌دستی، توجه به ظرفیت‌های "
          "گردشگری، انتخاب لوکیشن‌های تاریخی و فرهنگی و پخش پلی‌بک‌های مرتبط، "
          "تصویری متنوع و چندبعدی از استان مرکزی به مخاطبان ارائه شود.")
CONC_3 = ("ویژگی شاخص این تجربه، پیوند میان محتوای برنامه و فضای واقعی فرهنگی و "
          "تاریخی استان بود؛ به‌گونه‌ای که معرفی ظرفیت‌ها صرفاً به گفت‌وگو محدود "
          "نماند و از طریق حضور در لوکیشن‌های متناسب و نمایش فعالیت‌های بومی و "
          "هنری، ابعاد تصویری و عینی نیز پیدا کرد.")
CONC_4 = ("در مجموع، عملکرد برنامه در این دو روز در راستای اهداف طرح ایران‌جان، "
          "مبنی بر معرفی ظرفیت‌های استانی، بازتاب تنوع فرهنگی و هنری ایران و "
          "بهره‌گیری از رسانه بوده است.")


def star(text, faithful, both=True):
    if not faithful:
        return text
    return f"*{text}*" if both else f"*{text}"


def build(path, faithful):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    add_page_border(section)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(13)
    nrpr = normal.element.get_or_add_rPr()
    nrfonts = nrpr.find(qn("w:rFonts"))
    if nrfonts is None:
        nrfonts = OxmlElement("w:rFonts")
        nrpr.append(nrfonts)
    nrfonts.set(qn("w:cs"), CS_FONT)
    nrfonts.set(qn("w:ascii"), LATIN_FONT)
    nrfonts.set(qn("w:hAnsi"), LATIN_FONT)
    _reorder(nrpr, RPR_ORDER)

    doc.core_properties.title = "گزارش عملکرد برنامه «به خانه برمی‌گردیم» ـ استان مرکزی"
    doc.core_properties.subject = SUBJECT

    add_par(doc, star("بسمه تعالی", faithful), size=14, bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

    add_par(doc, star(TITLE, faithful), size=15, bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16, line=1.5)

    add_par(doc, star("مدیر محترم شبکه تهران سیما", faithful), size=13, bold=True,
            align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=14)

    add_par(doc, star(SUBJECT, faithful, both=False), size=13,
            align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=32)

    add_par(doc, star("اهداف حضور", faithful), size=14, bold=True,
            align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

    add_par(doc, INTRO_1)
    add_par(doc, INTRO_2, space_after=28)

    add_par(doc, star("دستاوردهای اجرایی و محتوایی", faithful), size=14, bold=True,
            align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

    add_par(doc, LEAD_IN, space_after=14)

    for num, lead, body in ITEMS:
        add_item(doc, num, lead, body, star_mark=faithful)

    doc.add_page_break()

    add_par(doc, star("جمع‌بندی و نتیجه‌گیری", faithful), size=14, bold=True,
            align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=14)

    add_par(doc, CONC_1)
    add_par(doc, CONC_2)
    add_par(doc, CONC_3)
    add_par(doc, CONC_4)

    body = doc.element.body
    for p in body.iter(qn("w:pPr")):
        _reorder(p, PPR_ORDER)
    for r in body.iter(qn("w:rPr")):
        _reorder(r, RPR_ORDER)
    for s in body.iter(qn("w:sectPr")):
        _reorder(s, SECT_ORDER)

    doc.save(path)
    print("saved:", path)


if __name__ == "__main__":
    base = "/home/user/nabz-sabz/"
    build(base + "گزارش عملکرد برنامه به خانه برمی‌گردیم - استان مرکزی.docx", True)
    build(base + "گزارش عملکرد برنامه به خانه برمی‌گردیم - استان مرکزی (بدون ستاره).docx", False)
