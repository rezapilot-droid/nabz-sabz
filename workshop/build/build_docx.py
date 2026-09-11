# -*- coding: utf-8 -*-
"""ساخت نسخهٔ Word (قابل ویرایش) کاربرگ — تک‌صفحه A4، راست‌به‌چپ"""
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(os.path.dirname(ROOT), "worksheets", "فصل۱-آزمایش۱")
os.makedirs(OUT_DIR, exist_ok=True)

# ---------------- رنگ‌ها
PURPLE  = RGBColor(0x8E, 0x44, 0xAD)
PINK    = RGBColor(0xD4, 0x54, 0x7E)
ROSE    = RGBColor(0xB2, 0x3E, 0x78)
BLUE    = RGBColor(0x3D, 0x7F, 0xA8)
GRAY    = RGBColor(0x8A, 0x75, 0x96)
BROWN   = RGBColor(0x9A, 0x6B, 0x12)
BODY    = RGBColor(0x54, 0x45, 0x5C)
GREEN   = RGBColor(0x1F, 0x8A, 0x63)
DOTCOL  = RGBColor(0xBB, 0x8F, 0xD4)

FA_FONT = "Vazirmatn"
FA_TITLE_FONT = "Lalezar"

def set_rtl_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def style_run(r, size, bold=False, color=BODY, font=FA_FONT):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font); rFonts.set(qn('w:hAnsi'), font); rFonts.set(qn('w:cs'), font)
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size*2))); rPr.append(szCs)
    if bold:
        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '1'); rPr.append(bCs)

def shade_para(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'), hex_fill)
    pPr.append(shd)

def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def table_rtl(tbl):
    tblPr = tbl._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual'); bidi.set(qn('w:val'),'1')
    tblPr.append(bidi)

def set_borders(tbl, color="F5B8D0", sz=12):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)

def para(cell_or_doc, first=False):
    if hasattr(cell_or_doc, 'paragraphs'):
        return cell_or_doc.paragraphs[0] if first else cell_or_doc.add_paragraph()
    return cell_or_doc.add_paragraph()

def sec_title(doc, emoji, text):
    p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{emoji} {text}")
    style_run(r, 12.5, True, RGBColor(0xFF,0xFF,0xFF))
    shade_para(p, "D989C4")
    return p

# ---------------- سند
doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.top_margin = Mm(9); sec.bottom_margin = Mm(9)
sec.left_margin = Mm(11); sec.right_margin = Mm(11)

st = doc.styles['Normal']
st.font.name = FA_FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:cs'), FA_FONT)

# هدر
p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("بسمه تعالی"); style_run(r, 12, True, PURPLE)
p.paragraph_format.space_after = Pt(4)

ttbl = doc.add_table(rows=1, cols=2); ttbl.alignment = WD_TABLE_ALIGNMENT.LEFT
table_rtl(ttbl)
cell_text = ttbl.rows[0].cells[0]; cell_img = ttbl.rows[0].cells[1]
cell_text.width = Mm(45); cell_img.width = Mm(22)
p = para(cell_text, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.RIGHT)
r = p.add_run("نام آموزگار:"); style_run(r, 9.5, False, RGBColor(0xA0, 0x6A, 0xB8))
p2 = cell_text.add_paragraph(); set_rtl_para(p2, WD_ALIGN_PARAGRAPH.RIGHT)
r = p2.add_run("خانم قاسم‌تبار"); style_run(r, 12.5, True, RGBColor(0x7C, 0x3F, 0x96))
pI = para(cell_img, True); set_rtl_para(pI, WD_ALIGN_PARAGRAPH.CENTER)
pI.add_run().add_picture(os.path.join(ROOT, "assets", "small", "teacher-avatar.png"), width=Mm(14))

p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("فصل اول: مخلوط و جداسازی مواد"); style_run(r, 17, True, PURPLE, FA_TITLE_FONT)
p.paragraph_format.space_after = Pt(1)

p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("آزمایش ۱: مخلوط‌های همگن و ناهمگن"); style_run(r, 13.5, True, PINK, FA_TITLE_FONT)
p.paragraph_format.space_after = Pt(1)

p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("برگرفته از فعالیت صفحهٔ ۳ کتاب علوم تجربی پایهٔ هشتم"); style_run(r, 8.5, False, GRAY)
p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.RIGHT)
r = p.add_run("نام و نام خانوادگی: .....................................     کلاس: هشتم .......     تاریخ: ....... / ....... / .......")
style_run(r, 9.5, False, BODY)
p.paragraph_format.space_after = Pt(2)

# --- لوازم و مواد
sec_title(doc, "🌸", "لوازم و مواد لازم برای آزمایش")
MATERIALS = [
    ("beaker-water.png",  "بشر شماره‌دار", "۲ عدد"),
    ("water-jug.png",     "آب", "۱۰۰ میلی‌لیتر در هر بشر"),
    ("measuring-spoon.png","قاشقک", "۱ عدد"),
    ("salt-pinches.png",  "نمک خوراکی", "۱ قاشقک سرپُر"),
    ("soil-spoon.png",    "خاک باغچه", "۱ قاشقک سرپُر"),
    ("rod-stir.png",      "همزن شیشه‌ای", "۱ عدد"),
]
tbl = doc.add_table(rows=2, cols=3); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
table_rtl(tbl); set_borders(tbl, "F6C6DD", 8)
for i,(img,name,amount) in enumerate(MATERIALS):
    cell = tbl.rows[i//3].cells[i%3]
    shade_cell(cell, "FFF6FB")
    cell.width = Mm(62)
    p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(2)
    p.add_run().add_picture(os.path.join(ROOT,"assets","small",img), height=Mm(21))
    p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(name); style_run(r, 10.5, True, ROSE)
    p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("مقدار: " + amount); style_run(r, 9, True, BLUE)
    p.paragraph_format.space_after = Pt(3)

# --- روش اجرا
sec_title(doc, "✨", "روش اجرای آزمایش")
STEPS = [
    "دو بشر را انتخاب کنید و با چسب کاغذی برای آن‌ها شمارهٔ ۱ و ۲ بگذارید.",
    "در هر دو بشر به مقدار یکسان (۱۰۰ میلی‌لیتر) آب بریزید.",
    "در بشر شمارهٔ ۱ یک قاشقک خاک باغچه و در بشر شمارهٔ ۲ یک قاشقک نمک خوراکی بریزید.",
    "محتویات هر دو بشر را با همزن شیشه‌ای به‌آرامی و کاملاً به هم بزنید.",
    "چند لحظه صبر کنید و سپس رنگ و شفافیت محتویات دو بشر را با هم مقایسه کنید.",
    "مشاهدات خود را در کادرهای بخش «مشاهدات و نتیجه‌گیری» یادداشت کنید.",
]
FA_NUM = "۱۲۳۴۵۶۷۸۹۰"
def fa(n): return "".join(FA_NUM[int(d)] for d in str(n))
for i, s in enumerate(STEPS, 1):
    p = doc.add_paragraph(); set_rtl_para(p)
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(1)
    r = p.add_run(f"{fa(i)}. "); style_run(r, 11, True, PINK)
    r = p.add_run(s); style_run(r, 10.5, False, BODY)

# --- مشاهدات
sec_title(doc, "💗", "مشاهدات و نتیجه‌گیری")
otbl = doc.add_table(rows=1, cols=2); table_rtl(otbl); set_borders(otbl, "ECC2E2", 8)
for j,(title,q,color,fill) in enumerate([
    ("بشر شمارهٔ ۱ — مخلوط آب و خاک","پس از هم زدن، مخلوط چگونه به نظر می‌رسد؟ (شفاف یا کدر؟)",ROSE,"FDF6FC"),
    ("بشر شمارهٔ ۲ — مخلوط آب و نمک","پس از هم زدن، مخلوط چگونه به نظر می‌رسد؟ (شفاف یا کدر؟)",GREEN,"F4FBF8"),
]):
    cell = otbl.rows[0].cells[j]; shade_cell(cell, fill); cell.width = Mm(94)
    p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(title); style_run(r, 10.5, True, color)
    p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(q); style_run(r, 8.5, False, GRAY)
    for _ in range(3):
        p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run("...............................................................")
        style_run(r, 9, False, DOTCOL)

p = doc.add_paragraph(); set_rtl_para(p)
p.paragraph_format.space_before = Pt(6)
r = p.add_run("نتیجه‌گیری: "); style_run(r, 10.5, True, PURPLE)
r = p.add_run("با توجه به مشاهدات، مخلوط شماره ..... همگن (محلول) و مخلوط شماره ..... ناهمگن است؛ زیرا")
style_run(r, 10.5, True, BODY)
p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("...................................................................................................................................")
style_run(r, 9, False, DOTCOL)

p = doc.add_paragraph(); set_rtl_para(p)
p.paragraph_format.space_before = Pt(6)
r = p.add_run("⚠ نکتهٔ ایمنی: "); style_run(r, 10, True, BROWN)
r = p.add_run("مواد آزمایشگاهی را نچشید و با چشمان خود به آن‌ها نزدیک نکنید؛ پس از پایان آزمایش، دست‌های خود را با آب و صابون بشویید.")
style_run(r, 9.5, False, BROWN)
shade_para(p, "FFF8E8")

out = os.path.join(OUT_DIR, "کاربرگ-فصل۱-آزمایش۱.docx")
doc.save(out)
print("DOCX saved:", out)
