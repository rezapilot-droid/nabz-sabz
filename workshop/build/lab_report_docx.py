# -*- coding: utf-8 -*-
"""قالب «گزارش آزمایشگاه علوم هشتم» — نسخهٔ Word قابل ویرایش (تک‌صفحه A4)"""
import json, os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEACHER_NAME = "خانم قاسم‌تبار"
HEADER_TITLE = "گزارش آزمایشگاه علوم هشتم"
FA_FONT = "Vazirmatn"
FA_TITLE_FONT = "Lalezar"
FA_NUM = "۰۱۲۳۴۵۶۷۸۹"
def fa(n): return "".join(FA_NUM[int(d)] for d in str(n))

PURPLE = RGBColor(0x8E, 0x44, 0xAD); PINK = RGBColor(0xD4, 0x54, 0x7E)
ROSE   = RGBColor(0xB2, 0x3E, 0x78); BLUE = RGBColor(0x3D, 0x7F, 0xA8)
GRAY   = RGBColor(0x8A, 0x75, 0x96); BROWN= RGBColor(0x9A, 0x6B, 0x12)
BODY   = RGBColor(0x54, 0x45, 0x5C); GREEN= RGBColor(0x0F, 0x7A, 0x55)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF); DOT  = RGBColor(0xBB, 0x8F, 0xD4)
HINTC  = RGBColor(0xA0, 0x5A, 0x92)

def set_rtl_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)

def style_run(r, size, bold=False, color=BODY, font=FA_FONT):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font); rFonts.set(qn('w:hAnsi'), font); rFonts.set(qn('w:cs'), font)
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size*2))); rPr.append(szCs)
    if bold:
        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '1'); rPr.append(bCs)

def shade_para(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'), hex_fill)
    pPr.append(shd)

def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def table_rtl(tbl):
    tblPr = tbl._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual'); bidi.set(qn('w:val'),'1'); tblPr.append(bidi)

def set_borders(tbl, color="F5B8D0", sz=8):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)

def para(container, first=False):
    if hasattr(container, 'paragraphs'):
        return container.paragraphs[0] if first else container.add_paragraph()
    return container.add_paragraph()

def sec_title(doc, num, text, is_ans=False):
    p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{fa(num)}. {text}")
    style_run(r, 11.5, True, WHITE)
    shade_para(p, "D989C4")

def dotted_lines(doc, n=2, size=10, color=DOT):
    for _ in range(n):
        p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("...................................................................................")
        style_run(r, size, False, color)

def build_docx(data, mode, out_path):
    is_ans = (mode == "answer")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(8); sec.bottom_margin = Mm(8)
    sec.left_margin = Mm(10); sec.right_margin = Mm(10)
    st = doc.styles['Normal']
    st.font.name = FA_FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:cs'), FA_FONT)

    # ---- هدر (جدول سه‌ستونه: نسخه | عنوان | دبیر)
    ht = doc.add_table(rows=1, cols=3); ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_rtl(ht)
    c_badge, c_title, c_teacher = ht.rows[0].cells
    for c, w in ((c_badge, Mm(30)), (c_title, Mm(96)), (c_teacher, Mm(52))):
        c.width = w
    shade_cell(c_badge, "E7F6EE" if is_ans else "E8F4FB")
    shade_cell(c_title, "D989C4")
    shade_cell(c_teacher, "FDE7F3")
    p = para(c_badge, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("پاسخنامه — ویژهٔ دبیر" if is_ans else "برگهٔ دانش‌آموز")
    style_run(r, 9.5, True, GREEN if is_ans else BLUE)
    p = para(c_title, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(HEADER_TITLE); style_run(r, 15, True, WHITE, FA_TITLE_FONT)
    p = para(c_teacher, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("دبیر آزمایشگاه: "); style_run(r, 9, False, HINTC)
    r = p.add_run(TEACHER_NAME); style_run(r, 10.5, True, RGBColor(0x7C, 0x3F, 0x96))

    p = doc.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(data['subject_short']); style_run(r, 13, True, PINK, FA_TITLE_FONT)

    # ---- مشخصات دانش‌آموز
    it = doc.add_table(rows=1, cols=3); table_rtl(it); set_borders(it, "E9B9D6", 6)
    labels = ["نام و نام خانوادگی: ...........................",
              "شماره کلاس: ..............",
              "تاریخ: ....... / ....... / ......."]
    for j, lab in enumerate(labels):
        cell = it.rows[0].cells[j]; shade_cell(cell, "FDF4FA")
        p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(lab); style_run(r, 9, True, RGBColor(0x8A, 0x5A, 0x9E))

    # ---- ۱. هدف (نسخهٔ دانش‌آموزی: خالی برای نوشتن)
    sec_title(doc, 1, "هدف آزمایش")
    if is_ans:
        p = doc.add_paragraph(); set_rtl_para(p); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(data['objective']); style_run(r, 9.5, False, BODY)
    else:
        dotted_lines(doc, 2)

    # ---- ۲. فرضیه
    sec_title(doc, 2, "فرضیه")
    if is_ans:
        p = doc.add_paragraph(); set_rtl_para(p); p.paragraph_format.space_after = Pt(2)
        r = p.add_run("پاسخ پیشنهادی: "); style_run(r, 9.5, True, GREEN)
        r = p.add_run(data['hypothesis']['answer']); style_run(r, 9.5, True, GREEN)
    else:
        dotted_lines(doc, 2)

    # ---- ۳. مواد و وسایل
    sec_title(doc, 3, "مواد و وسایل")
    mats = data['materials']
    tbl = doc.add_table(rows=2, cols=3); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_rtl(tbl); set_borders(tbl, "F6C6DD", 6)
    for i, m in enumerate(mats):
        cell = tbl.rows[i // 3].cells[i % 3]; shade_cell(cell, "FFF6FB")
        p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        img_path = (os.path.join(ROOT, "assets", m["img"]) if m["img"].startswith("icons/")
                    else os.path.join(ROOT, "assets", "small", m["img"]))
        p.add_run().add_picture(img_path, height=Mm(14))
        p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(m['name']); style_run(r, 9.5, True, ROSE)
        p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(m['amount']); style_run(r, 8.5, True, BLUE)
        p.paragraph_format.space_after = Pt(2)

    # ---- ۴. روش کار (نسخهٔ دانش‌آموزی: خالی برای نوشتن)
    sec_title(doc, 4, "روش کار")
    if is_ans:
        for i, step in enumerate(data['procedure'], 1):
            p = doc.add_paragraph(); set_rtl_para(p)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"{fa(i)}. "); style_run(r, 10, True, PINK)
            r = p.add_run(step); style_run(r, 9.5, False, BODY)
    else:
        dotted_lines(doc, 4)

    # ---- ۵. جدول ثبت مشاهدات
    sec_title(doc, 5, "جدول ثبت مشاهدات")
    obs = data['observation']
    ot = doc.add_table(rows=1 + len(obs['rows']), cols=len(obs['columns']))
    table_rtl(ot); set_borders(ot, "E5A9C9", 6)
    for j, c in enumerate(obs['columns']):
        cell = ot.rows[0].cells[j]; shade_cell(cell, "D989C4")
        p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(c); style_run(r, 9, True, WHITE)
    for i, row in enumerate(obs['rows'], 1):
        cell0 = ot.rows[i].cells[0]; shade_cell(cell0, "FDF0F8")
        p = para(cell0, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(row['label']); style_run(r, 9, True, ROSE)
        for j in range(1, len(obs['columns'])):
            cell = ot.rows[i].cells[j]
            p = para(cell, True); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
            if is_ans:
                r = p.add_run(row['answer'][j - 1]); style_run(r, 9, True, GREEN)
            else:
                r = p.add_run(" "); style_run(r, 9)

    # ---- ۶. نتیجه‌گیری
    sec_title(doc, 6, "نتیجه‌گیری")
    if is_ans:
        p = doc.add_paragraph(); set_rtl_para(p); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(data['conclusion']['answer']); style_run(r, 9.5, True, GREEN)
    else:
        dotted_lines(doc, 2)

    # ---- ۷. فکر کنید و پاسخ دهید
    sec_title(doc, 7, "فکر کنید و پاسخ دهید")
    tt = doc.add_table(rows=1, cols=2); table_rtl(tt); set_borders(tt, "ECC2E2", 6)
    for j, t in enumerate(data['think']):
        cell = tt.rows[0].cells[j]; shade_cell(cell, "FDF6FC")
        p = para(cell, True); set_rtl_para(p)
        r = p.add_run(f"{fa(j+1)}. {t['q']}"); style_run(r, 9, True, PURPLE)
        if is_ans:
            p = cell.add_paragraph(); set_rtl_para(p)
            r = p.add_run("پاسخ: " + t['answer']); style_run(r, 8.8, True, GREEN)
        else:
            for _ in range(2):
                p = cell.add_paragraph(); set_rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
                r = p.add_run("...............................................")
                style_run(r, 8.5, False, DOT)

    # ---- فوتر «طراحی شده با عشق»
    lp = doc.add_paragraph(); set_rtl_para(lp, WD_ALIGN_PARAGRAPH.CENTER)
    lp.paragraph_format.space_before = Pt(8)
    r = lp.add_run("🌸 طراحی شده با ❤ برای دانش‌آموزان دبیرستان هوردخت 🌸")
    style_run(r, 10.5, True, RGBColor(0xC2, 0x44, 0x7C))
    shade_para(lp, "FFEEF8")

    doc.save(out_path)
    print("DOCX saved:", out_path)

if __name__ == "__main__":
    import sys
    slug = sys.argv[1] if len(sys.argv) > 1 else "fasl1-azmayesh1"
    with open(os.path.join(ROOT, "experiments", slug + ".json"), encoding="utf-8") as f:
        data = json.load(f)
    for mode in ("student", "answer"):
        build_docx(data, mode, os.path.join(ROOT, "build", f"{slug}-{mode}.docx"))
